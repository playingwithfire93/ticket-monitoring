from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont
import os

CV_DOCX = 'CV_Tripulante_Espanol.docx'
CV_PDF = 'CV_Tripulante_Espanol.pdf'

def make_placeholder(path='photo_placeholder.png', size=(600,600)):
    img = Image.new('RGB', size, (220,220,220))
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype('arial.ttf', 36)
    except Exception:
        font = ImageFont.load_default()
    text = 'FOTO\nAQUÍ'
    try:
        bbox = draw.multiline_textbbox((0,0), text, font=font)
        w = bbox[2] - bbox[0]
        h = bbox[3] - bbox[1]
    except Exception:
        # fallback: estimate using textsize on single line
        w, h = draw.textsize(text, font=font)
    draw.multiline_text(((size[0]-w)/2,(size[1]-h)/2), text, fill=(80,80,80), font=font, align='center')
    img.save(path)
    return path

def add_heading(doc, name):
    p = doc.add_paragraph()
    run = p.add_run(name)
    run.bold = True
    run.font.size = Pt(18)
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

def add_contact_block(doc, email, phone, city, driving):
    p = doc.add_paragraph()
    p.add_run(f'Email: {email}  |  Tel: {phone}  |  {city}\n')
    p.add_run(f'Permiso: {driving}  |  Disponibilidad para viajar y reubicación')

def build_docx(output=CV_DOCX, photo_path=None):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Header: name and contact
    add_heading(doc, 'Blanca Téllez Valle')
    add_contact_block(doc, 'adatellezvalle@gmail.com', '602-502-302', 'Marbella', 'B y vehículo propio')

    # Create a 2-column table: left for text, right for photo placeholder
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(4.5)
    table.columns[1].width = Inches(1.8)

    left_cell = table.rows[0].cells[0]
    right_cell = table.rows[0].cells[1]

    # Left: profile and sections
    left_cell_par = left_cell.paragraphs[0]
    left_cell_par.add_run('Perfil profesional\n').bold = True
    left_cell_par.add_run('Tripulante orientada al servicio con amplia experiencia en atención al cliente, gestión de grupos y resolución de incidencias en entornos dinámicos. Mantengo la calma bajo presión, aplico rigurosamente protocolos de seguridad y garantizo una experiencia de viaje segura y cercana para el pasajero. Nivel de inglés C1, comunicación multicultural y alta capacidad de adaptación; motivada y comprometida con la excelencia en servicio a bordo.\n\n')

    left_cell.add_paragraph('Experiencia profesional').runs[0].bold = True
    left_cell.add_paragraph('Junior Automation Developer — Capgemini (Mar 2024 – Actualidad)')
    left_cell.add_paragraph('- Trabajo en entornos de alta exigencia y coordinación con equipos multidisciplinares.')
    left_cell.add_paragraph('- Gestión del tiempo, priorización y resolución de problemas bajo presión.\n')

    left_cell.add_paragraph('Profesora de Programación — Logiscool España (Sep 2022 – Jun 2024)')
    left_cell.add_paragraph('- Comunicación efectiva con alumnos y familias; atención personalizada y manejo de grupos.')
    left_cell.add_paragraph('- Empatía, paciencia y resolución de incidencias en tiempo real.\n')

    left_cell.add_paragraph('Guía de Zoo — Bioparc Fuengirola (Nov 2019 – Mar 2024)')
    left_cell.add_paragraph('- Atención a visitantes nacionales e internacionales; gestión de grupos guiados.')
    left_cell.add_paragraph('- Resolución de conflictos y cumplimiento de protocolos de seguridad.\n')

    left_cell.add_paragraph('Dependienta — ZARA SA (Ago 2022 – Sep 2022)')
    left_cell.add_paragraph('- Atención al cliente, ventas y trabajo en equipo en entorno de ritmo elevado.\n')

    left_cell.add_paragraph('Cajera / Reponedora — Carrefour (Jul 2019 – Ago 2019)')
    left_cell.add_paragraph('- Atención al cliente, gestión de incidencias y responsabilidad en caja.\n')

    left_cell.add_paragraph('Formación académica').runs[0].bold = True
    left_cell.add_paragraph('Grado en Ciencias Ambientales — Universidad de Málaga\n')

    left_cell.add_paragraph('Idiomas').runs[0].bold = True
    left_cell.add_paragraph('Español: Nativo  •  Inglés: Avanzado (C1)  •  Italiano: Básico\n')

    left_cell.add_paragraph('Competencias clave').runs[0].bold = True
    left_cell.add_paragraph('- Servicio al pasajero • Gestión de conflictos • Protocolos de seguridad\n- Trabajo bajo presión • Resolución de incidencias • Empatía y trato profesional\n')

    left_cell.add_paragraph('Información adicional').runs[0].bold = True
    left_cell.add_paragraph('Permiso de conducir B y vehículo propio. Disponibilidad para viajar y reubicación.').italic = True

    # Right: photo placeholder
    if photo_path and os.path.exists(photo_path):
        right_cell.add_paragraph().add_run().add_picture(photo_path, width=Inches(1.6))
    else:
        placeholder = make_placeholder()
        right_cell.add_paragraph().add_run().add_picture(placeholder, width=Inches(1.6))

    doc.save(output)
    print(f'DOCX guardado en: {output}')

    # Try to convert to PDF using docx2pdf if available
    try:
        from docx2pdf import convert
        convert(output, CV_PDF)
        print(f'PDF guardado en: {CV_PDF}')
    except Exception:
        print('\nNo se pudo generar el PDF automáticamente.\nOpciones:')
        print('- Instala `docx2pdf` y Microsoft Word en Windows, y ejecuta este script de nuevo. (pip install docx2pdf)')
        print('- O abre el .docx en Word y exporta a PDF manualmente.')

if __name__ == "__main__":
    # Si quieres usar tu propia foto, pásala como primer argumento: python generate_cv.py mi_foto.jpg
    import sys
    photo = sys.argv[1] if len(sys.argv) > 1 else None
    build_docx(photo_path=photo)
